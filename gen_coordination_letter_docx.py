from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

style = doc.styles['Normal']
style.font.name = '仿宋'
style.font.size = Pt(14)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
style.paragraph_format.line_spacing = 1.5

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '黑体'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hs.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        hs.font.size = Pt(18)
    elif level == 2:
        hs.font.size = Pt(15)
    elif level == 3:
        hs.font.size = Pt(14)

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(3.17)
section.right_margin = Cm(3.17)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)


def set_cell_shading(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.name = '黑体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        set_cell_shading(hdr_cells[i], 'D9E2F3')
    for ri, row_data in enumerate(rows):
        row_cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row_data):
            row_cells[ci].text = val
            for p in row_cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.name = '仿宋'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
    return table


def add_para(doc, text, bold=False, indent=False, font_size=14, font_name='仿宋', align=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(font_size * 2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run.element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if align:
        p.alignment = align
    return p


def add_mixed_para(doc, parts, indent=False, font_size=14):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(font_size * 2)
    for text, bold in parts:
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.name = '仿宋'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
    return p


# ==================== Title ====================
add_para(doc, '关于"河北好药数字化体系"GAP与GMP数据衔接工作的协调函', bold=True, font_size=18, font_name='黑体', align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()

# ==================== Header info ====================
add_para(doc, '致：河北好药中药饮片科技有限公司（业主）、GMP开发团队')
add_para(doc, '发件人：GAP开发项目组（码码通团队）')
add_para(doc, '日期：2026年5月13日')
add_para(doc, '主题：关于明确GAP与GMP数据衔接标准及协同要求的函')

doc.add_paragraph()

# ==================== Opening ====================
add_para(doc, '为确保"河北好药"项目顺利落地，实现从"田间地头"到"生产车间"的全链条数字化追溯，满足河北省药监局对GAP（种植）与GMP（加工）数据"无缝衔接、穿透核查"的监管要求，我方（GAP开发团队）就数据协同事项与贵方（GMP开发团队）沟通如下，请贵方确认并配合执行，以免因数据标准不一导致系统返工或合规风险。', indent=True)

doc.add_paragraph()

# ==================== Section 1 ====================
doc.add_heading('一、数据标准与定义统一要求', level=2)

add_para(doc, '为避免"同名异义"或"同物异名"导致的数据断层，双方应遵循统一的基础数据定义：', indent=True)

doc.add_heading('1.1 药材品种名称', level=3)
add_para(doc, '双方系统必须使用统一的标准药典名称（如"防风"而非"北防风"，"柴胡"而非"硬柴胡"）。我方将提供《康保道地药材品种编码表》，贵方系统需直接引用该编码表，避免自行创建新名称。', indent=True)

doc.add_heading('1.2 批次号映射规则', level=3)
add_table(doc, ['批次类型', '前缀', '示例', '生成方'], [
    ['采收批次号', 'CJ-', 'CJ-20260515-001', 'GAP端'],
    ['内部物料批号', 'YL-', 'YL-20260517-005', 'GMP端'],
    ['成品批号', 'CP-', 'CP-20260520-003', 'GMP端'],
])

add_mixed_para(doc, [('关联要求：', True), ('贵方在生成YL-批次号时，必须关联并存储对应的CJ-批次号，确保批次链条可追溯。建议采用以下关联规则：', False)], indent=True)

add_para(doc, '一个CJ-批次可拆分为多个YL-批次（如分批加工）', indent=True)
add_para(doc, '多个CJ-批次不建议合并为一个YL-批次（影响溯源精度）', indent=True)
add_para(doc, '如确需合并，需在系统中记录合并原因及审批人', indent=True)

doc.add_heading('1.3 计量单位统一', level=3)
add_table(doc, ['指标', '统一单位'], [
    ['重量', 'kg'],
    ['面积', '亩'],
    ['水分含量', '%'],
    ['温度', '°C'],
    ['湿度', '%RH'],
])

doc.add_paragraph()

# ==================== Section 2 ====================
doc.add_heading('二、接口对接与时间节点', level=2)

add_para(doc, '为保障"万码互通"，双方需配合完成以下接口联调工作：', indent=True)

add_table(doc, ['序号', '事项', '责任方', '截止时间', '具体要求'], [
    ['1', '接口规范确认', '双方', '2026年6月5日', '共同确认《GAP-GMP数据接口规范》，明确字段含义、调用方式、异常处理'],
    ['2', '测试环境联调', '双方', '2026年6月20日', '我方提供测试用采收批次数据（含地块、农事、初验数据），贵方完成数据拉取与解析测试'],
    ['3', '生产环境切换', '双方', '2026年7月10日', '关闭测试接口，正式启用生产环境数据同步'],
    ['4', '异常处理机制', '双方', '持续', '任何一方调用数据失败，需通过"码码通"企业微信预警群反馈，双方承诺4小时内响应'],
])

doc.add_paragraph()

# ==================== Section 3 ====================
doc.add_heading('三、数据质量与责任边界', level=2)

add_para(doc, '根据"谁产生、谁负责"原则，明确双方在数据链条中的责任：', indent=True)

doc.add_heading('3.1 GAP端责任（我方）', level=3)
add_para(doc, '确保源头数据（地块、农事、采收）的真实性（GPS定位、水印照片）与不可篡改性（审计追踪）', indent=True)
add_para(doc, '确保CJ-批次号在系统内唯一且状态正常（未被删除或作废）', indent=True)
add_para(doc, '在采收批次生成后24小时内，通过接口向贵方推送采收数据（含地块信息、农事摘要、初验结果）', indent=True)
add_para(doc, '如CJ-批次状态变更（如作废、修正），需在2小时内通知贵方', indent=True)

doc.add_heading('3.2 GMP端责任（贵方）', level=3)
add_para(doc, '在扫码入库（扫描CJ-码生成YL-码）时，校验我方提供的初验结果。若我方标记为"拒收"，贵方系统应提示拦截，禁止生成加工任务', indent=True)
add_para(doc, '在加工过程中若发现原料质量问题（如霉变、虫蛀），需通过接口即时回写质量异常信息至我方系统，以便我方更新地块档案', indent=True)
add_para(doc, '贵方人员不应在后台手动修改或伪造与我方CJ-批次的关联关系；如确需修正，需通过审批流程并记录审计日志', indent=True)

doc.add_heading('3.3 共同责任', level=3)
add_para(doc, '双方均需确保接口调用的数据完整性和及时性', indent=True)
add_para(doc, '双方均需保留完整的操作审计日志，以备药监局穿透式飞检', indent=True)
add_para(doc, '如发现数据异常，发现方需第一时间通知对方，双方共同排查', indent=True)

doc.add_paragraph()

# ==================== Section 4 ====================
doc.add_heading('四、现场实操与硬件协同', level=2)

add_para(doc, '为确保农户和车间工人操作顺畅，请在硬件部署上予以配合：', indent=True)

doc.add_heading('4.1 扫码设备兼容性', level=3)
add_para(doc, '我方生成的二维码均为标准QR Code，包含URL及明文批次号。请贵方确认车间PDA或工业平板支持扫描通用QR Code。', indent=True)

doc.add_heading('4.2 赋码打印协同', level=3)
add_para(doc, '我方输出的采收批次标签（CJ码）尺寸为50mm × 30mm。请贵方在规划标签打印机时，预留相应规格的纸仓空间，确保标签能被正常粘贴和扫描。', indent=True)

doc.add_heading('4.3 网络环境', level=3)
add_para(doc, '如车间网络不稳定，建议贵方在关键工位部署离线缓存机制，待网络恢复后自动同步数据。', indent=True)

doc.add_paragraph()

# ==================== Section 5 ====================
doc.add_heading('五、合规与审计配合', level=2)

add_para(doc, '鉴于药监局将实施"穿透式飞检"：', indent=True)

add_para(doc, '1. 请贵方在系统演示或迎检时，配合使用我方提供的"码码通"小程序作为数据源入口，共同展示从成品倒推至种植地块的全过程，以体现数据链条的真实性与连贯性。', indent=True)
add_para(doc, '2. 如遇飞检需调取种植端原始数据（如土壤报告、农事记录），我方将开放临时只读权限或导出PDF，请贵方提前熟悉数据调取路径。', indent=True)
add_para(doc, '3. 双方应定期（建议每月一次）进行数据一致性校验，确保GAP端与GMP端的批次关联关系准确无误。', indent=True)

doc.add_paragraph()

# ==================== Closing ====================
add_mixed_para(doc, [('以上建议旨在保障项目合规、高效推进。请贵方于', False), ('2026年5月25日前', True), ('书面确认回执。如有异议或补充意见，欢迎随时沟通协商。', False)], indent=True)

doc.add_paragraph()
add_para(doc, '顺祝商祺！', indent=True)

doc.add_paragraph()
doc.add_paragraph()

add_para(doc, '河北好药数字化体系GAP开发项目组', align=WD_ALIGN_PARAGRAPH.RIGHT)
add_para(doc, '2026年5月13日', align=WD_ALIGN_PARAGRAPH.RIGHT)

doc.add_paragraph()

# ==================== Receipt ====================
add_para(doc, '回执确认', bold=True, font_size=14, font_name='黑体')

add_para(doc, '□ 已收到协调函，同意上述要求，将按时配合执行。')
add_para(doc, '□ 已收到协调函，部分同意，具体意见如下：_________________________')
add_para(doc, '□ 有异议，具体意见如下：_________________________')

doc.add_paragraph()

add_para(doc, '确认单位（盖章）：_________________________')
add_para(doc, '确认人：_________________________')
add_para(doc, '日期：2026年____月____日')

import os
outpath = os.path.join('D:/code_cys/mtong', '关于\u201c河北好药数字化体系\u201dGAP与GMP数据衔接工作的协调函_修改.docx')
doc.save(outpath)
print(f'Saved: {outpath}')
